#!/usr/bin/env python3
"""
Fill in the `mice` table from the lab surgery workbook.

USAGE
    python import_mice.py MAIN_FOLDER SURGERY_XLSX [--mouse N]... [--project P]
                          [--dry-run]

Without `--mouse`, it does every animal that already has a session in the
database, which is the usual case: the recordings are ingested first and this
puts a name to them. Anything the workbook does not say is left alone, so it is
safe to rerun after the workbook is corrected or extended.

The workbook is keyed on SID, the colony's own number, which is what recent
folder names mean by `m442`. Older recordings used a separate numbering and
will not match; those are reported, not guessed at.

READS
    `line` as a cross of mutations, e.g. `TH-Cre x TIGRE`, and `genotype` as one
    word per mutation in that cross, e.g. `het het`. A line with no genotype
    still records the mutations, with the genotype left unknown; a genotype with
    the wrong number of words is reported and its genotypes dropped.

REQUIRES
    openpyxl to read the workbook.
"""

from __future__ import annotations

import argparse
import sys
from datetime import datetime
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from odyn import Database  # noqa: E402

# Searched in order, first hit wins. The surgery logs carry the line and the
# genotype; the colony lists are the fallback for an animal never operated on.
SHEETS = [
    "PA surgery log - Moss Lab",
    "Surgery log",
    "Available",
    "Moss Lab Mice - SIDs",
]

# The label each field goes by, across the two sheet layouts.
FIELDS = {
    "sid": ("SID", "Mouse SID"),
    "line": ("line", "Mouseline"),
    "genotype": ("genotype", "Genotype"),
    "sex": ("sex", "Sex"),
    "dob": ("DOB", "Date of Birth"),
}

# One spelling per mutation, so 'Dat-Cre' and 'DAT-Cre' are not two lines.
LINE_NAMES = {
    "dat-cre": "DAT-Cre",
    "th-cre": "TH-Cre",
    "pv-cre": "PV-Cre",
    "tigre": "TIGRE",
    "slc32a": "Slc32a",
    "thy1-gcamp": "Thy1-GCaMP",
    "thy1 gcamp": "Thy1-GCaMP",
    "thy1-gcamp6f": "Thy1-GCaMP6f",
    "thy1-gcamp8": "Thy1-GCaMP8",
}

# Written in the line column to mean the animal carries no mutation at all,
# which is the absence of a line rather than a line by that name.
NO_LINES = {"wt", "wild type", "wildtype", "-"}


def read_workbook(path: Path) -> dict[int, dict]:
    """Every animal the workbook knows, by SID."""
    book = openpyxl.load_workbook(path, data_only=True, read_only=True)
    found: dict[int, dict] = {}

    try:
        for sheet in SHEETS:
            if sheet not in book.sheetnames:
                continue

            for sid, record in read_sheet(book[sheet]).items():
                found.setdefault(sid, record)
    finally:
        book.close()

    return found


def read_sheet(worksheet) -> dict[int, dict]:
    """One sheet's rows, by SID, or nothing if it has no SID column."""
    rows = list(worksheet.iter_rows(values_only=True))
    index = header_index(rows)

    if index is None:
        return {}

    header_row, columns = index
    records = {}

    for row in rows[header_row + 1:]:
        sid = value(row, columns, "sid")

        if not isinstance(sid, (int, float)) or int(sid) <= 0:
            continue

        records.setdefault(int(sid), {
            name: value(row, columns, name) for name in FIELDS if name != "sid"
        })

    return records


def header_index(rows: list[tuple]) -> None | tuple[int, dict[str, int]]:
    """
    Which row holds the labels, and which column each field is in.

    The surgery sheets put a grouping row above the labels and the colony lists
    do not, so the header is found by looking for the SID column rather than
    assumed to be first.
    """
    for number, row in enumerate(rows[:4]):
        labels = {
            str(cell).replace("\n", " ").strip(): column
            for column, cell in enumerate(row)
            if cell is not None
        }

        columns = {
            name: next(
                (labels[label] for label in aliases if label in labels), None
            )
            for name, aliases in FIELDS.items()
        }

        if columns["sid"] is not None:
            return number, {k: v for k, v in columns.items() if v is not None}

    return None


def value(row: tuple, columns: dict[str, int], name: str):
    """The cell for one field, with blanks flattened to `None`."""
    column = columns.get(name)

    if column is None or column >= len(row):
        return None

    cell = row[column]

    return None if isinstance(cell, str) and not cell.strip() else cell


def parse_lines(line, genotype) -> tuple[dict[str, None | str], None | str]:
    """
    The mutations of a cross and their genotypes, plus a complaint if they clash.

    **EXAMPLE**
    ```python
    parse_lines("TH-Cre x TIGRE", "het het")
    # ({"TH-Cre": "het", "TIGRE": "het"}, None)
    ```
    """
    if line is None or str(line).strip().lower() in NO_LINES:
        return {}, None

    parts = [part.strip() for part in str(line).split(" x ") if part.strip()]
    parts = [LINE_NAMES.get(part.lower(), part) for part in parts]

    if not parts:
        return {}, None

    words = str(genotype).split() if genotype is not None else []

    # One word per mutation is the whole reason the two columns can be read
    # together. When they disagree, which mutation each word describes is a
    # guess, so the mutations are kept and the genotypes are not.
    if words and len(words) != len(parts):
        complaint = (
            f"line {line!r} has {len(parts)} parts but genotype {genotype!r} "
            f"has {len(words)} words; genotypes dropped"
        )

        return dict.fromkeys(parts), complaint

    # No genotype is not a problem: plenty of animals are recorded by line and
    # genotyped later, or never.
    return dict(zip(parts, words or [None] * len(parts))), None


def parse_dob(dob) -> None | str:
    """The date of birth as `YYYY-MM-DD`, or `None` if it is not a date."""
    if isinstance(dob, datetime):
        return dob.date().isoformat()

    if dob is None:
        return None

    # The colony lists write it as MM-DD-YYYY text rather than a real date.
    for layout in ("%m-%d-%Y", "%Y-%m-%d", "%m/%d/%Y"):
        try:
            return datetime.strptime(str(dob).strip(), layout).date().isoformat()
        except ValueError:
            continue

    return None


def parse_sex(sex) -> None | str:
    """`'M'` or `'F'`, from whichever of the spellings the sheet uses."""
    if sex is None:
        return None

    first = str(sex).strip()[:1].upper()

    return first if first in ("M", "F") else None


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("main_folder", type=Path)
    parser.add_argument("surgery_xlsx", type=Path)
    parser.add_argument("--mouse", action="append", type=int, dest="mice")
    parser.add_argument("--project", default=None)
    parser.add_argument("--dry-run", action="store_true")

    args = parser.parse_args()

    db = (
        Database(args.main_folder, project=args.project)
        if args.project
        else Database(args.main_folder)
    )

    wanted = args.mice or sorted(
        int(sid)
        for (sid,) in db.con.execute("SELECT DISTINCT mouse_id FROM sessions;")
    )

    if not wanted:
        print("No sessions in the database and no --mouse given, nothing to do.")
        return 0

    records = read_workbook(args.surgery_xlsx)
    missing = [sid for sid in wanted if sid not in records]

    for sid in wanted:
        record = records.get(sid)

        if record is None:
            continue

        carried, complaint = parse_lines(record["line"], record["genotype"])

        if complaint:
            print(f"  m{sid}: {complaint}")

        fields = {
            "sex": parse_sex(record["sex"]),
            "dob": parse_dob(record["dob"]),
            "lines": carried or None,
        }

        summary = " x ".join(
            f"{line}({genotype or '?'})" for line, genotype in carried.items()
        )

        print(
            f"  m{sid}: {fields['sex'] or '?'} "
            f"born {fields['dob'] or '?'} {summary or 'unknown line'}"
        )

        if not args.dry_run:
            db.set_mouse(mouse_id=sid, **fields)

    done = len(wanted) - len(missing)
    action = "would be written" if args.dry_run else "written"

    print(f"\n{done} of {len(wanted)} mice {action}.")

    if missing:
        print(
            f"Not in the workbook: {missing}. Recordings older than the SID "
            f"numbering use numbers that mean something else there."
        )

    return 0


if __name__ == "__main__":
    raise SystemExit(main())
